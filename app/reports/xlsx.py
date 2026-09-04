"""Workbook plumbing shared by every report builder.

Two things live here:

* :func:`to_xlsx_bytes` -- the `BytesIO` + `pd.ExcelWriter(engine='openpyxl')`
  + `to_excel(index=False)` (+ optional per-sheet formatter) dance that each
  export page repeated inline. Same engine, same argument order, same
  per-sheet formatter call after the write, so the bytes it returns are the
  bytes the pages produced.

* The six `_apply_price_formatting*` variants, moved VERBATIM from the page
  modules under distinct names. They are genuinely six different functions --
  different column lists, different number formats -- and are deliberately
  NOT merged: each one is the exact formatter one family of workbooks was
  written with, and merging them would silently move number formats between
  files.

Nothing in this module imports Streamlit or reads session state.
"""
from io import BytesIO

import pandas as pd


def to_xlsx_bytes(sheets: dict[str, pd.DataFrame], formatter=None) -> bytes:
    """Write `sheets` (sheet name -> DataFrame) to an .xlsx and return its bytes.

    Args:
        sheets: Ordered mapping of sheet name to DataFrame. Sheets are written
            in iteration order, each with `index=False`, exactly as the pages
            wrote them.
        formatter: Optional `f(writer, sheet_name, df)` applied after each
            sheet is written -- one of the `apply_*_price_formatting`
            functions below, or None for workbooks that were written unformatted.

    Returns:
        bytes: the workbook, ready for `st.download_button(data=...)`.
    """
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            if formatter is not None:
                formatter(writer, sheet_name, df)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# The six formatters, moved verbatim. Source module named above each one.
# ---------------------------------------------------------------------------

# was `app/pages/results/components/export_section.py::_apply_price_formatting`
def apply_export_price_formatting(writer, sheet_name: str, df: pd.DataFrame):
    """
    Apply Excel number formatting to numeric columns.
    
    Uses different decimal precision based on column type:
    - 4 decimal places for trait values (Agreeable, Openness, etc.)
    - 5 decimal places for intercept
    - 6 decimal places for calculated values (PB_i, DI_i)
    - 2 decimal places for prices, scores, and other values
    
    This formats the DISPLAY only - the underlying values retain full precision.
    
    Args:
        writer: ExcelWriter object
        sheet_name: Name of the sheet to format
        df: DataFrame being written (to identify column positions)
    """
    from openpyxl.styles import numbers
    
    # Columns with 4 decimal places (trait values from disclose_income)
    four_decimal_columns = [
        'Agreeable', 'Openness', 'Honesty_Humility', 'Extraversion', 'Neuroticism',
        'Religious', 'TWT+Sospeso', 'WOPB', 'WPB'
    ]
    
    # Columns with 5 decimal places (intercept)
    five_decimal_columns = ['Intercept']
    
    # Columns with 6 decimal places (calculated values)
    six_decimal_columns = ['PB_i', 'DI_i']
    
    # Define columns that should display with 2 decimal places
    two_decimal_columns = [
        # Vendor attributes
        'Vendor Price', 'avg_vendor_price', 'price',
        # Customer prices
        'Customer Price', 'customer_price',
        # Bid values
        'Bid Value', 'bid_value',
        # Scores (normalized 0-1, but show 2 decimals for clarity)
        'Standardized Vendor Price Score', 'Standardized Vendor Quality Score',
        'Standardized Vendor Sustainability Score', 'Standardized Vendor Proximity Score',
        'Vendor Integrated Score', 'avg_vendor_score',
        # Donation rates
        'donation_default', 'Agent Donation Default', 'Final Donation Rate',
        'final_donation_rate', 'avg_vendor_quality', 'avg_vendor_sustainability',
        'avg_vendor_proximity', 'Vendor Proximity', 'Vendor Quality', 'Vendor Sustainability',
        # Income
        'income', 'Income',
        # Vendor choice weights
        'weight_price', 'weight_quality', 'weight_proximity', 'weight_sustainability',
    ]
    
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    # Get column indices and apply appropriate formatting
    for col_idx, col_name in enumerate(df.columns, start=1):
        # Determine format based on column name
        if col_name in six_decimal_columns:
            number_format = '0.000000'
        elif col_name in five_decimal_columns:
            number_format = '0.00000'
        elif col_name in four_decimal_columns:
            number_format = '0.0000'
        elif col_name in two_decimal_columns:
            number_format = '0.00'
        else:
            continue  # Skip columns not in any list
        
        # Apply number format to entire column (skip header row)
        for row_idx in range(2, len(df) + 2):  # Start from row 2 (after header)
            cell = worksheet.cell(row=row_idx, column=col_idx)
            if isinstance(cell.value, (int, float)) and cell.value is not None:
                cell.number_format = number_format


# was `app/pages/results/visualizations/disclosure_viz.py::_apply_price_formatting_disclosure`
def apply_disclosure_price_formatting(writer, sheet_name: str, df: pd.DataFrame):
    """
    Apply Excel number formatting to numeric columns.
    
    Uses 'General' format to preserve original decimal precision from the data.
    No rounding or truncation is applied - values display exactly as stored.
    """
    # All numeric columns use General format to preserve original precision
    numeric_columns = [
        'Agreeable', 'Openness', 'Honesty_Humility', 'Extraversion', 'Neuroticism',
        'Religious', 'TWT+Sospeso', 'calc_PB', 'WOPB', 'WPB', 'Intercept', 'PB_i', 'Disclosure Income',
        'income', 'Income', 'TWT+Sospeso [=AW2+AX2]{Periods 1+2}',
        'Assigned income from the distribution',
        # Disclose Documents (item-22 layout)
        'PersonalIncentive', 'PrivacyConcern', 'Trust', 'Disclosure Document',
    ]
    
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name not in numeric_columns:
            continue
        
        # Apply General format to preserve original precision
        for row_idx in range(2, len(df) + 2):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            if isinstance(cell.value, (int, float)) and cell.value is not None:
                cell.number_format = 'General'


# was `app/pages/results/visualizations/transaction_viz.py::_apply_price_formatting_transaction`
def apply_transaction_price_formatting(writer, sheet_name: str, df: pd.DataFrame):
    """
    Apply Excel number formatting to price-related columns to display 2 decimal places.
    """
    price_columns = [
        'Customer Price', 'customer_price', 'Bid Value', 'bid_value',
        'Final Donation Rate', 'final_donation_rate',
    ]
    
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name in price_columns:
            for row_idx in range(2, len(df) + 2):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)) and cell.value is not None:
                    cell.number_format = '0.00'


# was `app/pages/results/visualizations/bidding_viz.py::_apply_price_formatting_bid`
def apply_bid_price_formatting(writer, sheet_name: str, df: pd.DataFrame):
    """
    Apply Excel number formatting to price-related columns to display 2 decimal places.
    """
    from openpyxl.styles import numbers
    
    price_columns = [
        'Honesty_Humility', 'income', 'Vendor Price', 'Bid Value',
        'TWT+Sospeso [=AW2+AX2]{Periods 1+2}'
    ]
    
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name in price_columns:
            for row_idx in range(2, len(df) + 2):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)) and cell.value is not None:
                    cell.number_format = '0.00'


# was `app/pages/results/visualizations/vendor_viz.py::_apply_price_formatting_vendor`
def apply_vendor_price_formatting(writer, sheet_name: str, df: pd.DataFrame):
    """
    Apply Excel number formatting to price-related columns to display 2 decimal places.
    """
    price_columns = [
        'Vendor Price', 'price', 'Customer Paid Price', 'Customer Price',
        'Bid Value', 'bid_value', 'Vendor Score', 'Integrated Score',
        'Price Score', 'Quality Score', 'Sustainability Score', 'Proximity Score',
        'Price (Normalized)', 'Quality (Normalized)', 'Sustainability (Normalized)', 'Proximity (Normalized)',
        'Weight_Price', 'Weight_Quality', 'Weight_Proximity', 'Weight_Sustainability',
        'weight_price', 'weight_quality', 'weight_proximity', 'weight_sustainability',
        'Proximity', 'Quality', 'Sustainability',
    ]
    
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name in price_columns:
            for row_idx in range(2, len(df) + 2):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)) and cell.value is not None:
                    cell.number_format = '0.00'


# was `app/pages/results/visualizations/donation_viz.py::_apply_price_formatting_donation`
def apply_donation_price_formatting(writer, sheet_name: str, df: pd.DataFrame):
    """
    Apply Excel number formatting to price-related columns to display 2 decimal places.
    
    This formats the DISPLAY only - the underlying values retain full precision.
    """
    from openpyxl.styles import numbers
    
    # Define columns that should display with 2 decimal places
    # Note: 'Donation Paid' and 'Total Paid by Customer' columns were removed
    # because actual price is unknown for Fixed/Discount customers
    price_columns = [
        'Customer Price', 'customer_price',
        'Default Donation Rate', 'Final Donation Rate',
        'donation_default', 'final_donation_rate',
        'Honesty_Humility', 'income', 'Income',
    ]
    
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    # Get column indices for price columns
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name in price_columns:
            # Apply number format to entire column (skip header row)
            for row_idx in range(2, len(df) + 2):  # Start from row 2 (after header)
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)) and cell.value is not None:
                    cell.number_format = '0.00'
